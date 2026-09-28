from __future__ import annotations
import csv, json, re
from datetime import date
from pathlib import Path
from typing import Any
from .base import MarketDataProvider

SUPPORTED = {'.csv','.xls','.xlsx','.json'}
ALIASES = {
 'listing_id':['listing_id','listing','listingid'], 'security_id':['security_id','security','securityid'],
 'ticker':['ticker','code','symbol','stock_code','证券代码','股票代码'], 'exchange':['exchange','exchange_mic','market','交易所'],
 'trade_date':['trade_date','date','tradedate','交易日期','日期'], 'open':['open','open_price','开盘'],
 'high':['high','high_price','最高'], 'low':['low','low_price','最低'], 'close':['close','close_price','收盘'],
 'adjusted_close':['adjusted_close','adj_close','后复权收盘'], 'volume':['volume','vol','成交量'],
 'turnover':['turnover','amount','turnover_value','成交额','成交金额'], 'currency':['currency','币种'],
 'source_id':['source_id','provider']
}

def _key(s): return re.sub(r'[^a-z0-9\u4e00-\u9fff]','',str(s).strip().lower())

def _map_headers(headers):
    m={_key(h):h for h in headers}
    out={}
    for target,names in ALIASES.items():
        for n in names:
            if _key(n) in m: out[target]=m[_key(n)]; break
    return out

def _read_rows(path: Path):
    ext=path.suffix.lower()
    if ext=='.csv':
        with path.open('r',encoding='utf-8-sig',newline='') as f: return list(csv.DictReader(f))
    if ext=='.json':
        obj=json.loads(path.read_text(encoding='utf-8'))
        if isinstance(obj,dict):
            for k in ('records','raw_rows','data','rows'):
                if isinstance(obj.get(k),list): return obj[k]
        if isinstance(obj,list): return obj
        raise ValueError('JSON does not contain a record list')
    if ext=='.xlsx':
        from openpyxl import load_workbook
        wb=load_workbook(path,data_only=True,read_only=True); ws=wb.active
        rows=list(ws.iter_rows(values_only=True));
        for i,row in enumerate(rows):
            hs=[str(x).strip() if x is not None else '' for x in row]
            if _map_headers(hs).get('trade_date') and (_map_headers(hs).get('close') or _map_headers(hs).get('open')):
                headers=hs; return [dict(zip(headers,[str(x).strip() if x is not None else None for x in r])) for r in rows[i+1:] if any(x is not None for x in r)]
        raise ValueError('header row not detected')
    if ext=='.xls':
        try:
            import xlrd
            wb=xlrd.open_workbook(path, on_demand=True); ws=wb.sheet_by_index(0); rows=[ws.row_values(i) for i in range(ws.nrows)]
            for i,row in enumerate(rows):
                hs=[str(x).strip() for x in row]
                if _map_headers(hs).get('trade_date') and (_map_headers(hs).get('close') or _map_headers(hs).get('open')):
                    headers=hs; return [dict(zip(headers,r)) for r in rows[i+1:] if any(x not in ('',None) for x in r)]
        except ImportError:
            raise ValueError('XLS parser requires xlrd; install xlrd or provide XLSX/CSV/JSON')
        raise ValueError('header row not detected')
    raise ValueError(f'unsupported format: {ext}')

def _float(v):
    if v is None or str(v).strip()=='': return None
    return float(str(v).replace(',',''))

def _ticker(v):
    s=str(v).strip().upper()
    m=re.search(r'(\d{6})(?:\.(?:SH|SZ))?$',s)
    return m.group(1) if m else s.zfill(6) if s.isdigit() else s

def _exchange(ticker, supplied=None):
    if supplied:
        s=str(supplied).strip().upper()
        aliases={'SH':'XSHG','SSE':'XSHG','XSHG':'XSHG','SZ':'XSHE','SZSE':'XSHE','XSHE':'XSHE'}
        if s in aliases: return aliases[s]
    if ticker.startswith(('600','601','603','605','688')): return 'XSHG'
    if ticker.startswith(('000','001','002','003','300','301')): return 'XSHE'
    return None

class ManualFileProvider(MarketDataProvider):
    source_origin='MANUAL_FILE'
    def __init__(self, path: Path, source_id: str='SRC-MANUAL-FILE'):
        self.path=Path(path); self.source_id=source_id
    def discover(self): return self.path.exists() and self.path.suffix.lower() in SUPPORTED
    def fetch(self): return self.path
    def parse(self,input_path=None): return _read_rows(input_path or self.path)
    def normalize(self,records):
        out=[]
        for r in records:
            h=_map_headers(r.keys())
            def get(k): return r.get(h[k]) if k in h else None
            ticker=_ticker(get('ticker'))
            ex=_exchange(ticker,get('exchange'))
            out.append({'listing_id':get('listing_id'),'security_id':get('security_id'),'ticker':ticker,'exchange_mic':ex,'normalized_ticker':f'CN-{ex}-{ticker}' if ex else None,'trade_date':str(get('trade_date'))[:10] if get('trade_date') is not None else None,'open':_float(get('open')),'high':_float(get('high')),'low':_float(get('low')),'close':_float(get('close')),'adjusted_close':_float(get('adjusted_close')),'price_type':'ADJUSTED_CLOSE' if get('adjusted_close') is not None else 'RAW_CLOSE','volume':_float(get('volume')),'turnover':_float(get('turnover')),'currency':get('currency') or 'CNY','source_id':get('source_id') or self.source_id,'source_origin':self.source_origin})
        return out
    def capabilities(self): return {'source_id':self.source_id,'dataset':'stock_daily','market':'CN','delivery_model':'MANUAL_DOWNLOAD','batch':True,'historical':True,'daily':True,'adjusted_price':'source_dependent','volume':True,'turnover':'source_dependent','market_cap':False,'valuation':False,'capital_flow':False,'access_status':'AVAILABLE','license_gate':'REVIEW','publication_scope':'INTERNAL_SNAPSHOT_ONLY'}
