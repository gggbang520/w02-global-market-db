from __future__ import annotations
from pathlib import Path
import re, shutil, subprocess, tempfile
from typing import Any

HEADER_ALIASES={
 'code':{'代码','证券代码','成分券代码','stock code','code','证券代码（含市场）'},
 'name':{'名称','证券简称','成分券名称','name'},
 'weight':{'权重','权重(%)','权重%','weight','weight(%)'},
 'exchange':{'交易所','市场','exchange','market'}
}
def normalize_code(v:Any)->str:
 s=str(v).strip();
 if s.endswith('.0') and s[:-2].isdigit(): s=s[:-2]
 m=re.search(r'(\d{6})',s)
 if not m: raise ValueError(f'invalid security code: {v!r}')
 return m.group(1)
def normalize_name(v:Any)->str: return re.sub(r'\s+','',str(v or '').strip())
def parse_weight(v:Any):
 if v is None or str(v).strip()=='': return None
 s=str(v).strip().replace('%','').replace(',',''); return float(s)
def map_exchange(code:str, explicit:Any=None)->str:
 e=str(explicit or '').upper().strip()
 if e in {'SSE','SH','XSHG','沪市','上海证券交易所'}: return 'XSHG'
 if e in {'SZSE','SZ','XSHE','深市','深圳证券交易所'}: return 'XSHE'
 return 'XSHG' if code.startswith(('600','601','603','605','688','689')) else 'XSHE'
def _header_map(row):
 out={}
 for i,v in enumerate(row):
  key=normalize_name(v).lower()
  for logical,aliases in HEADER_ALIASES.items():
   if key in {normalize_name(a).lower() for a in aliases}: out[logical]=i
 return out
def _rows_from_xlsx(path:Path):
 from openpyxl import load_workbook
 wb=load_workbook(path,read_only=True,data_only=True); ws=wb.active; rows=list(ws.iter_rows(values_only=True)); wb.close(); return rows
def _rows_from_xls(path:Path):
 try:
  import xlrd
  book=xlrd.open_workbook(path.as_posix(),on_demand=True); sh=book.sheet_by_index(0); rows=[sh.row_values(i) for i in range(sh.nrows)]; book.release_resources(); return rows
 except ImportError:
  if not shutil.which('libreoffice'): raise RuntimeError('XLS parser unavailable: xlrd and libreoffice are both missing')
  with tempfile.TemporaryDirectory(prefix='w02_xls_') as td:
   out=Path(td); p=subprocess.run(['libreoffice','--headless','--convert-to','xlsx','--outdir',str(out),str(path)],capture_output=True,text=True,timeout=60); converted=out/(path.stem+'.xlsx')
   if p.returncode!=0 or not converted.exists(): raise RuntimeError(f'XLS conversion failed: {p.stderr.strip() or p.stdout.strip()}')
   return _rows_from_xlsx(converted)

def parse_csi_file_detailed(path:str|Path,source_origin:str='FIXTURE')->dict:
 path=Path(path); rows=_rows_from_xlsx(path) if path.suffix.lower()=='.xlsx' else _rows_from_xls(path) if path.suffix.lower()=='.xls' else (_ for _ in ()).throw(ValueError(f'unsupported file extension: {path.suffix}'))
 header_idx=None; hm=None
 for i,row in enumerate(rows[:100]):
  candidate=_header_map(row)
  if 'code' in candidate and 'name' in candidate: header_idx=i; hm=candidate; break
 if header_idx is None: raise ValueError('SCHEMA_CHANGED: code/name headers not detected')
 raw_rows=max(0,len(rows)-header_idx-1); parsed=[]; invalid=0; missing_code=missing_name=missing_weight=0
 for row in rows[header_idx+1:]:
  if not row or all(v is None or str(v).strip()=='' for v in row): continue
  try: code=normalize_code(row[hm['code']])
  except Exception: invalid+=1; missing_code+=1; continue
  name=normalize_name(row[hm['name']]) if hm['name']<len(row) else ''
  if not name: invalid+=1; missing_name+=1; continue
  weight=None
  if 'weight' in hm and hm['weight']<len(row):
   try: weight=parse_weight(row[hm['weight']])
   except Exception: invalid+=1; continue
  if weight is None: missing_weight+=1
  explicit=row[hm['exchange']] if 'exchange' in hm and hm['exchange']<len(row) else None
  exchange=map_exchange(code,explicit)
  parsed.append({'ticker':code,'source_ticker':str(row[hm['code']]).strip(),'name':name,'weight':weight,'exchange_mic':exchange,'source_origin':source_origin,'effective_from':None,'effective_to':None})
 codes=[r['ticker'] for r in parsed]; dup=sorted({c for c in codes if codes.count(c)>1})
 validation={'header_row':header_idx,'raw_rows':raw_rows,'parsed_rows':len(parsed),'valid_rows':len(parsed)-len(dup),'invalid_rows':invalid,'duplicate_rows':len(dup),'duplicate_tickers':dup,'missing_code':missing_code,'missing_name':missing_name,'missing_weight':missing_weight,'source_origin':source_origin}
 return {'records':parsed,'validation':validation}

def parse_csi_file(path:str|Path,source_origin:str='FIXTURE')->list[dict]:
 d=parse_csi_file_detailed(path,source_origin); 
 if d['validation']['duplicate_rows']: raise ValueError(f"duplicate security codes: {d['validation']['duplicate_tickers'][:10]}")
 return d['records']
